"""项目管理（PM）Excel 导出接口（CH-009）。

覆盖看板与四个子页，全部复用 `app.exporter` 的样式与工具，保证与报价模块导出的
观感一致（同一套表头填充、边框、等宽字体列宽）。

| 页面 | 接口 | 内容 |
|---|---|---|
| 进度看板 | `GET /api/pm/board/export` | 多项目进度清单（日期/进度/剩余/工序完成数） |
| 工期计划表 | `GET /api/pm/projects/{id}/stages/export` | 该项目 8 阶段 + 合计行 |
| 施工实施步骤 | `GET /api/pm/processes/export?project_id=` | 工序 + 状态 + 负责人 + 实际开始/完成 |
| 预算决算管理点 | `GET /api/pm/cost-points/export?project_id=` | 环节/管理点/控制动作/风险/责任主体 |
| 预算决算配套单据 | `GET /api/pm/documents/export?project_id=` | 审批节点/当前节点/归档文件完成情况 |

权限：`require_internal`；`project_id` 不传 = 标准模板，仅管理员可导出（与列表接口一致）。
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import require_internal
from ..exporter import (
    _FILL_HEAD,
    _FILL_TOTAL,
    _dt,
    _safe,
    _today,
    _write_header,
    _write_row,
    content_disposition,
)
from ..models import PmCostPoint, PmDocument, PmProcess, PmProject, User, app_now
from .pm import _ensure_project_exists, _guard_template, _load_project
from .pm import _progress as _pm_progress
from .pm_documents import document_out

router = APIRouter()

STATUS_LABELS = {
    "planning": "未开工",
    "active": "进行中",
    "done": "已竣工",
    "paused": "暂停",
}
PROCESS_STATUS_LABELS = {
    "todo": "未开始",
    "doing": "进行中",
    "done": "已完成",
    "paused": "暂停",
}


def _xlsx(wb: Workbook, filename: str, ascii_name: str) -> Response:
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(filename, ascii_name)},
    )


def _widths(ws, widths: tuple[int, ...]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _progress(project: PmProject, today: date) -> tuple[int, int | None]:
    """复用列表接口的口径（CH-012：按当前阶段；日期只算剩余天数），避免两套算法漂移。"""
    info = _pm_progress(project, today)
    return int(info["progress"]), info["remaining_days"]  # type: ignore[arg-type]


# --------------------------------------------------------------------------
# 看板：多项目进度清单
# --------------------------------------------------------------------------
@router.get("/board/export", summary="导出多项目进度清单（xlsx）")
def export_board(
    db: Session = Depends(get_db),
    user: User = Depends(require_internal),
):
    today = app_now().date()
    projects = list(db.scalars(select(PmProject).order_by(PmProject.id)).all())

    # 各项目工序完成数
    counts: dict[int, tuple[int, int]] = {}
    for project_id, status, count in db.execute(
        select(PmProcess.project_id, PmProcess.status, func.count())
        .where(PmProcess.project_id.is_not(None))
        .group_by(PmProcess.project_id, PmProcess.status)
    ).all():
        done, total = counts.get(project_id, (0, 0))
        total += count
        if status == "done":
            done += count
        counts[project_id] = (done, total)

    wb = Workbook()
    ws = wb.active
    ws.title = "多项目进度"
    headers = [
        "序号", "项目名称", "项目编号", "当前阶段", "状态", "项目经理",
        "开工日期", "计划竣工", "总工期(月)", "总体进度%", "剩余天数",
        "工序完成", "备注",
    ]
    _write_header(ws, headers)
    for index, project in enumerate(projects, start=1):
        pct, remaining = _progress(project, today)
        done, total = counts.get(project.id, (0, 0))
        _write_row(
            ws,
            [
                index,
                project.name,
                project.code or "",
                project.current_stage or "",
                STATUS_LABELS.get(project.status, project.status),
                project.manager or "",
                project.start_date.isoformat() if project.start_date else "",
                project.plan_end_date.isoformat() if project.plan_end_date else "",
                float(project.total_months) if project.total_months is not None else "",
                pct,
                remaining if remaining is not None else "",
                f"{done}/{total}",
                project.remark or "",
            ],
        )
    _widths(ws, (6, 28, 16, 12, 10, 12, 12, 12, 12, 12, 12, 10, 24))
    ws.freeze_panes = "A2"
    return _xlsx(
        wb,
        f"多项目进度_{_today().strftime('%Y%m%d')}.xlsx",
        f"pm-projects-{_today().strftime('%Y%m%d')}.xlsx",
    )


# --------------------------------------------------------------------------
# 工期计划表
# --------------------------------------------------------------------------
@router.get("/projects/{project_id}/stages/export", summary="导出项目工期计划表（xlsx）")
def export_project_stages(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_internal),
):
    project = _load_project(db, project_id)
    stages = sorted(project.stages, key=lambda s: s.seq)

    wb = Workbook()
    ws = wb.active
    ws.title = "工期计划表"
    headers = [
        "序号", "阶段名称", "阶段组", "工期(月)", "开始月", "结束月",
        "关键交付物 / 里程碑", "项目经理角色", "责任主体",
    ]
    _write_header(ws, headers)
    build_total = 0.0
    pm_total = 0.0
    for stage in stages:
        months = float(stage.months) if stage.months is not None else 0.0
        if stage.seq <= 7:
            build_total += months
        if 4 <= stage.seq <= 7:
            pm_total += months
        _write_row(
            ws,
            [
                stage.seq,
                stage.name,
                stage.group_name or "",
                months,
                float(stage.start_month) if stage.start_month is not None else "",
                float(stage.end_month) if stage.end_month is not None else "",
                stage.deliverables or "",
                stage.pm_role or "",
                stage.owner or "",
            ],
        )
    all_total = build_total + float(
        next((s.months for s in stages if s.seq == 8), 0) or 0
    )
    total_row = _write_row(
        ws,
        ["合计", f"建设工期（1~7）{build_total:g} 月", f"项目经理在场（4~7）{pm_total:g} 月",
         build_total, "", "", f"含质保合计 {all_total:g} 月", "", ""],
        bold=True,
        fill=_FILL_TOTAL,
    )
    ws.freeze_panes = "A2"
    _widths(ws, (6, 16, 12, 10, 9, 9, 40, 18, 14))
    ws.append([])
    _write_row(
        ws,
        [f"项目：{project.name}",
         f"状态：{STATUS_LABELS.get(project.status, project.status)}",
         f"开工：{project.start_date or '—'}",
         f"计划竣工：{project.plan_end_date or '—'}"],
        bold=True,
        fill=_FILL_HEAD,
    )
    return _xlsx(
        wb,
        f"工期计划表_{_safe(project.name)}_{_today().strftime('%Y%m%d')}.xlsx",
        f"pm-stages-{project.id}-{_today().strftime('%Y%m%d')}.xlsx",
    )


# --------------------------------------------------------------------------
# 三张参考表
# --------------------------------------------------------------------------
@router.get("/processes/export", summary="导出施工实施步骤明细（xlsx）")
def export_processes(
    project_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_internal),
):
    scope = "标准模板"
    if project_id:
        project = _ensure_project_exists(db, project_id)
        db.commit()
        scope = project.name
        rows = db.scalars(
            select(PmProcess).where(PmProcess.project_id == project_id).order_by(PmProcess.seq)
        ).all()
    else:
        _guard_template(user)
        rows = db.scalars(
            select(PmProcess).where(PmProcess.project_id.is_(None)).order_by(PmProcess.seq)
        ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "施工实施步骤明细"
    headers = [
        "序号", "工序分组", "工序名称", "主要工作内容", "质量控制点",
        "紧前工序", "负责人", "当前状态", "实际开始", "实际完成", "最近进度说明",
    ]
    _write_header(ws, headers)
    for row in rows:
        _write_row(
            ws,
            [
                row.seq,
                row.group_name,
                row.name,
                row.content or "",
                row.qc_point or "",
                row.predecessor or "",
                row.assignee or "",
                PROCESS_STATUS_LABELS.get(row.status, row.status),
                _dt(row.started_at),
                _dt(row.finished_at),
                row.progress_remark or "",
            ],
        )
    _widths(ws, (6, 12, 22, 40, 32, 14, 12, 10, 18, 18, 28))
    ws.freeze_panes = "A2"
    return _xlsx(
        wb,
        f"施工实施步骤明细_{_safe(scope)}_{_today().strftime('%Y%m%d')}.xlsx",
        f"pm-processes-{project_id or 'template'}-{_today().strftime('%Y%m%d')}.xlsx",
    )


@router.get("/cost-points/export", summary="导出预算决算管理点（xlsx）")
def export_cost_points(
    project_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_internal),
):
    scope = "标准模板"
    if project_id:
        project = _ensure_project_exists(db, project_id)
        db.commit()
        scope = project.name
        rows = db.scalars(
            select(PmCostPoint).where(PmCostPoint.project_id == project_id).order_by(PmCostPoint.seq)
        ).all()
    else:
        _guard_template(user)
        rows = db.scalars(
            select(PmCostPoint).where(PmCostPoint.project_id.is_(None)).order_by(PmCostPoint.seq)
        ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "预算决算管理点"
    headers = ["序号", "环节", "管理点", "关键控制动作", "常见风险 / 红线", "责任主体"]
    _write_header(ws, headers)
    for row in rows:
        _write_row(
            ws,
            [row.seq, row.phase, row.point, row.action or "", row.risk or "", row.owner or ""],
        )
    _widths(ws, (6, 12, 20, 44, 36, 14))
    ws.freeze_panes = "A2"
    return _xlsx(
        wb,
        f"预算决算管理点_{_safe(scope)}_{_today().strftime('%Y%m%d')}.xlsx",
        f"pm-cost-points-{project_id or 'template'}-{_today().strftime('%Y%m%d')}.xlsx",
    )


@router.get("/documents/export", summary="导出预算决算配套单据（xlsx）")
def export_documents(
    project_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_internal),
):
    scope = "标准模板"
    if project_id:
        project = _ensure_project_exists(db, project_id)
        db.commit()
        scope = project.name
        rows = db.scalars(
            select(PmDocument).where(PmDocument.project_id == project_id).order_by(PmDocument.seq)
        ).all()
    else:
        _guard_template(user)
        rows = db.scalars(
            select(PmDocument).where(PmDocument.project_id.is_(None)).order_by(PmDocument.seq)
        ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "预算决算配套单据"
    headers = [
        "序号", "单据名称", "对应管理点", "核心栏目", "审批流转节点",
        "当前节点", "归档文件", "已归档数", "是否完成归档", "归档完成时间", "归档人",
    ]
    _write_header(ws, headers)
    for row in rows:
        out = document_out(row)
        files = "；".join(
            f"{'✓' if f.archived else '○'} {f.name}{f'（{f.note}）' if f.note else ''}"
            for f in out.archive_files
        )
        _write_row(
            ws,
            [
                out.seq,
                out.name,
                out.cost_point or "",
                out.required_fields or "",
                " → ".join(out.flow_nodes),
                out.current_node or "",
                files,
                f"{out.archive_done}/{out.archive_total}",
                "是" if out.archived else "否",
                _dt(out.archived_at),
                out.archived_by or "",
            ],
        )
    _widths(ws, (6, 22, 18, 34, 30, 14, 42, 10, 12, 18, 12))
    ws.freeze_panes = "A2"
    return _xlsx(
        wb,
        f"预算决算配套单据_{_safe(scope)}_{_today().strftime('%Y%m%d')}.xlsx",
        f"pm-documents-{project_id or 'template'}-{_today().strftime('%Y%m%d')}.xlsx",
    )
