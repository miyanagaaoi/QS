"""Excel 导出：导入模板、询价单导出、比价明细导出、汇总统计导出（PRD 5.1.3 / 5.1.5 / 5.3.2）。

统一约定：
- 价格列套用人民币两位小数格式 `"￥"#,##0.00`（FR-EXP-04）
- 最低价单元格底色高亮
- 文件名规范 `询价单_{单号}_{项目}_{YYYYMMDD}.xlsx`（FR-EXP-03）
- 单元格写入 Decimal 本身（openpyxl 支持 Decimal），全程不转 float
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import app_now
from .schemas import ComparisonResult, SummaryResult
from .utils import to_app_tz

STATUS_LABEL = {"draft": "草稿", "published": "已发布", "closed": "已截止"}
ASSIGN_LABEL = {"pending": "待报价", "partial": "报价中", "quoted": "已报价"}

EXPORT_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MONEY_FMT = '"￥"#,##0.00'
QTY_FMT = "#,##0.###"

_FONT_TITLE = Font(bold=True, size=14)
_FONT_BOLD = Font(bold=True)
_FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")
_FILL_BEST = PatternFill("solid", fgColor="FFF2CC")
_FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
_SIDE = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")


# --------------------------------------------------------------------------
# 通用工具
# --------------------------------------------------------------------------
def _safe(text: str | None) -> str:
    """文件名安全化：去掉 Windows 非法字符与空白。"""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t\s]+', "", text or "")
    return cleaned or "未命名"


def _dt(value: datetime | None) -> str:
    """应用基准（北京时间）naive -> 可读字符串（内网单时区，不做时区换算）。"""
    value = to_app_tz(value)
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _today() -> date:
    """导出文件名里的「当天」必须用应用基准（北京时间），不能用宿主机的 date.today()。

    PRD 部署方案是 Linux 容器（通常 UTC）：若用宿主本地日期，北京时间 00:00~08:00
    导出的文件会带上"昨天"的日期，与界面展示的截止/创建时间不一致。
    """
    return app_now().date()


def rfq_export_filename(rfq_no: str, project: str, day: date | None = None) -> str:
    return f"询价单_{_safe(rfq_no)}_{_safe(project)}_{(day or _today()).strftime('%Y%m%d')}.xlsx"


def comparison_export_filename(rfq_no: str, project: str, day: date | None = None) -> str:
    return f"比价明细_{_safe(rfq_no)}_{_safe(project)}_{(day or _today()).strftime('%Y%m%d')}.xlsx"


def summary_export_filename(day: date | None = None) -> str:
    return f"报价单统计_{(day or _today()).strftime('%Y%m%d')}.xlsx"


def content_disposition(filename: str, ascii_fallback: str) -> str:
    """同时给出 ASCII 兜底名与 UTF-8 文件名（中文名在浏览器可正常下载）。"""
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"


def _decimal(value: str | None) -> Decimal | None:
    if value in (None, ""):
        return None
    return Decimal(str(value))


def _write_row(ws, values: list, *, bold: bool = False, fill: PatternFill | None = None) -> int:
    ws.append(values)
    row_no = ws.max_row
    for col in range(1, len(values) + 1):
        cell = ws.cell(row=row_no, column=col)
        cell.border = _BORDER
        if bold:
            cell.font = _FONT_BOLD
        if fill is not None:
            cell.fill = fill
    return row_no


def _write_header(ws, headers: list[str]) -> int:
    row_no = _write_row(ws, headers, bold=True, fill=_FILL_HEAD)
    for col in range(1, len(headers) + 1):
        ws.cell(row=row_no, column=col).alignment = _CENTER
    return row_no


# --------------------------------------------------------------------------
# 导入模板（FR-IMP-01）
# --------------------------------------------------------------------------
def build_import_template(units: list[str]) -> bytes:
    """生成导入模板：明细 Sheet（6 列）+ 填写说明 Sheet。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "询价明细"
    headers = ["序号", "材料名称", "规格型号", "单位", "项目", "采购数量"]
    _write_header(ws, headers)
    ws.append([1, "订书机", "得力，小", "个", "办公用品", 40])
    for col, width in zip(range(1, 7), (8, 24, 24, 10, 18, 12)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    if units:
        # 单位下拉（数据有效性），仅作录入便利，后端仍以字典为准
        joined = ",".join(units)[:250]
        dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True, showDropDown=False)
        dv.error = "单位必须在计量单位字典内"
        dv.prompt = "从下拉中选择计量单位"
        ws.add_data_validation(dv)
        dv.add(f"D2:D{2001}")

    guide = wb.create_sheet("填写说明")
    rows = [
        ["列名", "是否必填", "填写说明"],
        ["序号", "否", "行内顺序号，不填则按行自动编号"],
        ["材料名称", "是", "最多 100 字，示例：订书机"],
        ["规格型号", "否", "最多 200 字，示例：得力，小"],
        ["单位", "是", "必须取自计量单位字典，最多 20 字，示例：个 / 台 / 套"],
        ["项目", "否", "行级项目名，最多 100 字，留空则继承询价单表头项目"],
        ["采购数量", "是", "大于 0 的数字，最多 3 位小数，示例：40"],
        ["", "", ""],
        ["注意事项", "", "1. 请勿修改表头文字与列顺序（备注列可选，如存在则一并导入）"],
        ["", "", "2. 单次导入上限 2000 行，超出请分批导入"],
        ["", "", "3. 上传后系统先做预校验并展示错误清单，确认后才写入询价单"],
        ["", "", "4. 单位不在字典内时，可在确认导入时选择跳过该行或自动加入字典"],
    ]
    for row in rows:
        guide.append(row)
    for col, width in zip(range(1, 4), (16, 10, 70)):
        guide.column_dimensions[get_column_letter(col)].width = width
    for cell in guide[1]:
        cell.font = _FONT_BOLD
        cell.fill = _FILL_HEAD

    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


# --------------------------------------------------------------------------
# 比价明细表（询价单导出与报表视图 A 共用）
# --------------------------------------------------------------------------
def _write_comparison_sheet(
    ws, result: ComparisonResult, *, info_rows: list[tuple[str, str]] | None = None, with_quotes: bool = True
) -> int:
    """写入比价明细表，返回表头行号。with_quotes=False 时只导出我方需求信息。"""
    if info_rows:
        ws.append([f"询价单 {result.rfq_no}"])
        ws.cell(row=ws.max_row, column=1).font = _FONT_TITLE
        for key, value in info_rows:
            ws.append([key, value])
            ws.cell(row=ws.max_row, column=1).font = _FONT_BOLD
            ws.cell(row=ws.max_row, column=1).alignment = _LEFT
        ws.append([])

    headers = ["序号", "材料名称", "规格型号", "单位", "项目", "采购数量"]
    supplier_ids = [s.supplier_id for s in result.suppliers] if with_quotes else []
    if with_quotes:
        headers += [s.supplier_name for s in result.suppliers]
        headers += ["最低价", "最低价供货商", "最高价", "均价", "价差率", "报价家数"]
    header_row = _write_header(ws, headers)

    for row in result.rows:
        values: list = [
            row.seq,
            row.material_name,
            row.spec or "",
            row.unit,
            row.project or "",
            Decimal(str(row.quantity)),
        ]
        best_col: int | None = None
        if with_quotes:
            for offset, supplier_id in enumerate(supplier_ids):
                price = _decimal(row.prices.get(str(supplier_id)))
                values.append(price if price is not None else "—")
                if price is not None and row.best_supplier_id == supplier_id:
                    best_col = 7 + offset
            values += [
                _decimal(row.best_price),
                row.best_supplier_name or "",
                _decimal(row.max_price),
                _decimal(row.avg_price),
                row.diff_rate or "",
                f"{row.quoted_count}/{row.required_count}",
            ]
        data_row = _write_row(ws, values)
        ws.cell(row=data_row, column=6).number_format = QTY_FMT
        if with_quotes:
            for col in range(7, 7 + len(supplier_ids)):
                ws.cell(row=data_row, column=col).number_format = MONEY_FMT
            money_start = 7 + len(supplier_ids)
            for col in range(money_start, money_start + 4):
                ws.cell(row=data_row, column=col).number_format = MONEY_FMT
        if best_col is not None:
            ws.cell(row=data_row, column=best_col).fill = _FILL_BEST
            ws.cell(row=data_row, column=best_col).font = _FONT_BOLD

    if with_quotes:
        total_row = ws.max_row + 1
        ws.cell(row=total_row, column=1, value="合计").font = _FONT_BOLD
        money_cell = ws.cell(row=total_row, column=money_start, value=_decimal(result.total_best_amount))
        money_cell.number_format = MONEY_FMT
        money_cell.font = _FONT_BOLD
        money_cell.fill = _FILL_TOTAL
        labels = [
            ("最优总金额", result.total_best_amount, True),
            ("最高总金额", result.total_max_amount, True),
            ("可节省金额", result.saved_amount, True),
            ("节省率", result.saved_rate, False),
        ]
        for offset, (label, value, is_money) in enumerate(labels):
            ws.cell(row=total_row + offset, column=money_start + 1, value=label).font = _FONT_BOLD
            cell = ws.cell(
                row=total_row + offset,
                column=money_start + 2,
                value=_decimal(value) if (is_money and value) else (value or ""),
            )
            if is_money:
                cell.number_format = MONEY_FMT

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 12
    for col in range(7, 7 + len(supplier_ids) + 5):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return header_row


def build_comparison_export(result: ComparisonResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "比价明细"
    _write_comparison_sheet(ws, result)

    sheet = wb.create_sheet("供货商")
    _write_header(sheet, ["供货商", "报价状态"])
    for supplier in result.suppliers:
        _write_row(sheet, [supplier.supplier_name, ASSIGN_LABEL.get(supplier.status, supplier.status)])
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 12
    return _save(wb)


# --------------------------------------------------------------------------
# 询价单导出（FR-EXP-01 / FR-EXP-02 / FR-EXP-04）
# --------------------------------------------------------------------------
def build_rfq_export(
    *,
    rfq_no: str,
    title: str,
    project: str,
    deadline: datetime | None,
    status: str,
    creator_name: str | None,
    created_at: datetime | None,
    remark: str | None,
    internal_remark: str | None,
    result: ComparisonResult,
    include_quotes: bool = True,
    progress: list[dict] | None = None,
) -> bytes:
    """导出询价单 xlsx：表头信息区 + 明细（+ 各家报价对比）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "询价单"
    info_rows = [
        ("单据名称", title),
        ("项目", project),
        ("报价截止时间", _dt(deadline)),
        ("状态", STATUS_LABEL.get(status, status)),
        ("创建人", creator_name or ""),
        ("创建时间", _dt(created_at)),
        ("备注", remark or ""),
        ("内部备注", internal_remark or ""),
        ("导出范围", "明细 + 各家报价对比" if include_quotes else "仅我方询价明细"),
    ]
    export_result = result if include_quotes else ComparisonResult(
        rfq_id=result.rfq_id,
        rfq_no=result.rfq_no,
        title=result.title,
        project=result.project,
        status=result.status,
        suppliers=[],
        rows=result.rows,
    )
    _write_comparison_sheet(
        ws,
        export_result,
        info_rows=[("询价单号", rfq_no)] + info_rows,
        with_quotes=include_quotes,
    )

    if include_quotes and progress:
        sheet = wb.create_sheet("报价进度")
        _write_header(sheet, ["供货商", "报价状态", "报价行数", "应填行数", "最后保存时间", "提交时间"])
        for item in progress:
            _write_row(
                sheet,
                [
                    item.get("supplier_name", ""),
                    ASSIGN_LABEL.get(item.get("status", ""), item.get("status", "")),
                    item.get("quoted_line_count", 0),
                    item.get("line_count", 0),
                    _dt(item.get("last_saved_at")),
                    _dt(item.get("submitted_at")),
                ],
            )
        for col, width in zip(range(1, 7), (24, 12, 12, 12, 22, 22)):
            sheet.column_dimensions[get_column_letter(col)].width = width
    return _save(wb)


# --------------------------------------------------------------------------
# 汇总统计导出（视图 B）
# --------------------------------------------------------------------------
def build_summary_export(result: SummaryResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "总览"
    _write_row(ws, ["报价单汇总统计"], bold=True)
    overview = [
        ("统计时间范围", f"{result.date_from or '不限'} ~ {result.date_to or '不限'}"),
        ("询价单数", result.rfq_count),
        ("明细行数", result.line_count),
        ("受邀家次", result.invited_assignments),
        ("已报价家次", result.quoted_assignments),
        ("响应率", result.response_rate or "0.0000"),
        ("最优总金额", _decimal(result.best_amount)),
    ]
    for label, value in overview:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = _FONT_BOLD
        ws.cell(row=ws.max_row, column=1).border = _BORDER
        cell = ws.cell(row=ws.max_row, column=2)
        cell.border = _BORDER
        if label == "最优总金额":
            cell.number_format = MONEY_FMT
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30

    supplier_sheet = wb.create_sheet("按供货商")
    _write_header(
        supplier_sheet,
        ["供货商", "受邀单数", "已报价单数", "响应率", "报价总金额", "平均交期(天)", "平均MOQ"],
    )
    for row in result.by_supplier:
        _write_row(
            supplier_sheet,
            [
                row.supplier_name,
                row.invited_count,
                row.quoted_count,
                row.response_rate or "0.0000",
                _decimal(row.amount),
                _decimal(row.avg_lead_time) if row.avg_lead_time else "",
                _decimal(row.avg_moq) if row.avg_moq else "",
            ],
        )
        cell = supplier_sheet.cell(row=supplier_sheet.max_row, column=5)
        cell.number_format = MONEY_FMT
    for col, width in zip(range(1, 8), (24, 12, 12, 12, 18, 14, 12)):
        supplier_sheet.column_dimensions[get_column_letter(col)].width = width

    project_sheet = wb.create_sheet("按项目")
    _write_header(project_sheet, ["项目", "询价单数", "明细行数", "报价覆盖率", "最优金额"])
    for row in result.by_project:
        _write_row(
            project_sheet,
            [
                row.project,
                row.rfq_count,
                row.line_count,
                row.coverage_rate or "0.0000",
                _decimal(row.best_amount),
            ],
        )
        project_sheet.cell(row=project_sheet.max_row, column=5).number_format = MONEY_FMT
    for col, width in zip(range(1, 6), (24, 12, 12, 14, 18)):
        project_sheet.column_dimensions[get_column_letter(col)].width = width

    return _save(wb)
