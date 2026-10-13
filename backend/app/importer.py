"""Excel 询价明细导入：解析、预校验、一次性令牌（PRD FR-IMP-01 ~ FR-IMP-08）。

设计要点：
- 支持 .xlsx（openpyxl）与 .xls（xlrd）。`doc/示例.xls` 虽为 OLE2 容器（伪装后缀），
  xlrd 2.0.1 实测可读，因此两种格式都直接解析，不依赖 LibreOffice 转换。
- 预校验**只读文件、不落库**；校验通过后签发一次性 token（内存缓存 + TTL），
  确认导入时按 token 取出解析结果落库（PRD FR-IMP-02）。
- 校验规则（FR-IMP-03）：必填缺失 / 数量非数字或 <= 0 / 单位不在字典 / 文本超长 / 整行空白。
- 错误精确到行列（FR-IMP-04），ImportErrorItem(row, column, value, reason)。
- 单位不在字典内属于「软错误」：仍签发 token，确认时由 add_unknown_units 决定
  跳过该行或自动加入字典（FR-IMP-07）。
"""
from __future__ import annotations

import io
import re
import secrets
import threading
import time
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

import xlrd
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import settings
from .models import UnitDict
from .schemas import ImportErrorItem, ImportPreviewLine, ImportPreviewResult

# 令牌有效期（秒）
TOKEN_TTL_SECONDS = 30 * 60
# 表头扫描范围（前 5 行内找表头，兼容文件顶部有标题行的情况）
HEADER_SCAN_ROWS = 5

SUPPORTED_EXT = (".xlsx", ".xls")


class ImportValidationError(Exception):
    """文件级错误（格式不支持 / 表头缺失 / 超行数上限），直接转 40001。"""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


# --------------------------------------------------------------------------
# 表头定义（对齐 5.1.2 明细行前 6 列 + 备注）
# --------------------------------------------------------------------------
# (字段名, 表头别名, 文本最大长度, 是否必填)
COLUMN_SPEC: tuple[tuple[str, tuple[str, ...], int | None, bool], ...] = (
    ("seq", ("序号",), None, False),
    ("material_name", ("材料名称", "名称"), 100, True),
    ("spec", ("规格型号", "规格"), 200, False),
    ("unit", ("单位",), 20, True),
    ("project", ("项目",), 100, False),
    ("quantity", ("采购数量", "数量"), None, True),
    ("remark", ("备注",), 200, False),
)


@dataclass
class ParsedRow:
    """一行解析结果（含校验标记）。"""

    row: int  # Excel 行号（1 基，含表头）
    seq: int
    material_name: str = ""
    spec: str = ""
    unit: str = ""
    project: str = ""
    quantity: Decimal | None = None
    remark: str = ""
    errors: list[ImportErrorItem] = field(default_factory=list)
    unknown_unit: bool = False

    @property
    def blocking(self) -> bool:
        """是否存在硬错误（硬错误阻断确认导入；未知单位不算硬错误）。"""
        for error in self.errors:
            is_unknown_unit = error.column == "单位" and "不在计量单位字典" in error.reason
            if not is_unknown_unit:
                return True
        return False

    def to_line(self) -> ImportPreviewLine:
        return ImportPreviewLine(
            seq=self.seq,
            material_name=self.material_name,
            spec=self.spec or None,
            unit=self.unit,
            project=self.project or None,
            quantity=format(self.quantity.quantize(Decimal("0.001")), "f") if self.quantity is not None else "0",
        )

    def to_payload(self) -> dict[str, Any]:
        """落库用字典（quantity 以字符串传递，调用侧再转 Decimal）。"""
        return {
            "seq": self.seq,
            "material_name": self.material_name,
            "spec": self.spec or None,
            "unit": self.unit,
            "project": self.project or None,
            "quantity": format(self.quantity.quantize(Decimal("0.001")), "f") if self.quantity is not None else "0",
            "remark": self.remark or None,
        }


@dataclass
class ImportToken:
    """一次性令牌缓存记录。"""

    token: str
    user_id: int
    rfq_id: int
    rows: list[ParsedRow]
    created_at: float


_TOKEN_STORE: dict[str, ImportToken] = {}
_TOKEN_LOCK = threading.Lock()


# --------------------------------------------------------------------------
# 单元格工具
# --------------------------------------------------------------------------
def _cell_text(value: Any) -> str:
    """单元格 -> 去空白字符串。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # 40.0 -> "40"；40.5 -> "40.5"
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    return str(value).strip()


def _norm_header(value: Any) -> str:
    """表头归一化：去空白、全角括号转半角。"""
    text = _cell_text(value)
    return re.sub(r"\s+", "", text).replace("（", "(").replace("）", ")")


def _parse_quantity(text: str) -> Decimal | None:
    """数量解析：非法或 <= 0 返回 None。"""
    if not text:
        return None
    try:
        value = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if value <= 0:
        return None
    return value


# --------------------------------------------------------------------------
# 文件解析
# --------------------------------------------------------------------------
def _read_xlsx(content: bytes) -> list[tuple[int, list[Any]]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        return [(idx, list(row)) for idx, row in enumerate(ws.iter_rows(values_only=True), start=1)]
    finally:
        wb.close()


def _read_xls(content: bytes) -> list[tuple[int, list[Any]]]:
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    return [
        (r + 1, [sheet.cell_value(r, c) for c in range(sheet.ncols)])
        for r in range(sheet.nrows)
    ]


def read_raw_rows(filename: str, content: bytes) -> list[tuple[int, list[Any]]]:
    """按扩展名解析，失败时按另一种格式回退（伪装的 .xls 也能读）。"""
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext not in SUPPORTED_EXT:
        raise ImportValidationError("仅支持 .xlsx 与 .xls 文件，请先另存为 Excel 格式")

    order = (_read_xlsx, _read_xls) if ext == ".xlsx" else (_read_xls, _read_xlsx)
    last_exc: Exception | None = None
    for reader in order:
        try:
            rows = reader(content)
            if rows:
                return rows
            last_exc = ImportValidationError("文件中没有任何内容")
        except Exception as exc:  # noqa: BLE001 - 统一转为可读错误
            last_exc = exc
    raise ImportValidationError(f"文件解析失败：{last_exc}")


def _detect_header(rows: list[tuple[int, list[Any]]]) -> tuple[int, dict[str, int]]:
    """定位表头行并建立「字段 -> 列下标」映射。"""
    for row_no, values in rows[:HEADER_SCAN_ROWS]:
        norm = [_norm_header(v) for v in values]
        mapping: dict[str, int] = {}
        for field_name, aliases, _max_len, _required in COLUMN_SPEC:
            for idx, cell in enumerate(norm):
                if cell and any(alias == cell or alias in cell for alias in aliases):
                    mapping[field_name] = idx
                    break
        if all(k in mapping for k in ("material_name", "unit", "quantity")):
            return row_no, mapping
    raise ImportValidationError("未找到表头列「材料名称 / 单位 / 采购数量」，请下载导入模板后按模板填写")


# --------------------------------------------------------------------------
# 预校验
# --------------------------------------------------------------------------
def parse_and_validate(db: Session, filename: str, content: bytes) -> list[ParsedRow]:
    """解析 + 逐行校验，返回全部数据行（含错误标记）。文件级错误抛 ImportValidationError。"""
    rows = read_raw_rows(filename, content)
    header_row, mapping = _detect_header(rows)

    enabled_units = {
        name for (name,) in db.execute(select(UnitDict.name).where(UnitDict.status == 1)).all()
    }

    data_rows = [(no, vals) for no, vals in rows if no > header_row]
    # 去掉尾部整行空白（Excel 常见的多余空行，不视为用户填写的内容）
    while data_rows and all(not _cell_text(v) for v in data_rows[-1][1]):
        data_rows.pop()

    if len(data_rows) > settings.import_max_rows:
        raise ImportValidationError(
            f"本次共 {len(data_rows)} 行，超过单次导入上限 {settings.import_max_rows} 行，请分批导入"
        )
    if not data_rows:
        raise ImportValidationError("文件中没有可导入的明细行")

    parsed: list[ParsedRow] = []
    for order, (row_no, values) in enumerate(data_rows, start=1):
        def cell(field_name: str) -> str:
            idx = mapping.get(field_name)
            if idx is None or idx >= len(values):
                return ""
            return _cell_text(values[idx])

        raw_values = [cell(f) for f, _a, _m, _r in COLUMN_SPEC]
        if not any(raw_values):
            parsed.append(
                ParsedRow(
                    row=row_no,
                    seq=order,
                    errors=[ImportErrorItem(row=row_no, column=None, reason="整行内容为空，请删除该行或补齐必填项")],
                )
            )
            continue

        item = ParsedRow(row=row_no, seq=order)
        seq_text = cell("seq")
        try:
            item.seq = int(Decimal(seq_text)) if seq_text else order
        except (InvalidOperation, ValueError):
            item.seq = order

        for field_name, aliases, max_len, required in COLUMN_SPEC:
            if field_name in ("seq",):
                continue
            text = cell(field_name)
            if field_name == "quantity":
                value = _parse_quantity(text)
                label = aliases[0]
                if not text:
                    if required:
                        item.errors.append(
                            ImportErrorItem(row=row_no, column=label, value="", reason="必填项缺失")
                        )
                elif value is None:
                    item.errors.append(
                        ImportErrorItem(
                            row=row_no, column=label, value=text, reason="必须为大于 0 的数字"
                        )
                    )
                else:
                    item.quantity = value
                continue
            if required and not text:
                item.errors.append(
                    ImportErrorItem(row=row_no, column=aliases[0], value="", reason="必填项缺失")
                )
            if max_len and len(text) > max_len:
                item.errors.append(
                    ImportErrorItem(
                        row=row_no,
                        column=aliases[0],
                        value=text[:30] + ("…" if len(text) > 30 else ""),
                        reason=f"文本超长（最多 {max_len} 字，当前 {len(text)} 字）",
                    )
                )
            setattr(item, field_name, text)

        if item.unit and item.unit not in enabled_units:
            item.unknown_unit = True
            item.errors.append(
                ImportErrorItem(
                    row=row_no,
                    column="单位",
                    value=item.unit,
                    reason="单位不在计量单位字典内（可在确认导入时选择「自动加入字典」，或先在字典中维护）",
                )
            )
        parsed.append(item)
    return parsed


def build_preview(db: Session, filename: str, content: bytes, *, user_id: int, rfq_id: int) -> ImportPreviewResult:
    """预校验并签发 token（无硬错误时）。"""
    parsed = parse_and_validate(db, filename, content)

    hard_error_rows = [r for r in parsed if r.blocking]
    ok_rows = [r for r in parsed if not r.errors]
    unknown_rows = [r for r in parsed if r.errors and not r.blocking]

    token: str | None = None
    if not hard_error_rows and (ok_rows or unknown_rows):
        token = secrets.token_urlsafe(24)
        with _TOKEN_LOCK:
            _purge_expired()
            _TOKEN_STORE[token] = ImportToken(
                token=token,
                user_id=user_id,
                rfq_id=rfq_id,
                rows=[r for r in parsed if not r.blocking],
                created_at=time.time(),
            )

    unknown_units = sorted({r.unit for r in parsed if r.unknown_unit and r.unit})
    errors: list[ImportErrorItem] = []
    for r in parsed:
        errors.extend(r.errors)

    return ImportPreviewResult(
        total_rows=len(parsed),
        valid_rows=len(ok_rows),
        error_rows=len(parsed) - len(ok_rows),
        preview=[r.to_line() for r in ok_rows[:20]],
        errors=errors,
        unknown_units=unknown_units,
        token=token,
    )


def _purge_expired() -> None:
    now = time.time()
    for key in [k for k, v in _TOKEN_STORE.items() if now - v.created_at > TOKEN_TTL_SECONDS]:
        _TOKEN_STORE.pop(key, None)


def take_token(token: str, user_id: int) -> ImportToken:
    """取出并作废 token（一次性）。越权/过期/不存在均抛 ImportValidationError。"""
    with _TOKEN_LOCK:
        _purge_expired()
        record = _TOKEN_STORE.get(token)
        if record is None:
            raise ImportValidationError("导入令牌无效或已过期，请重新上传文件预校验")
        if record.user_id != user_id:
            raise ImportValidationError("导入令牌不属于当前用户，请重新上传文件预校验")
        _TOKEN_STORE.pop(token, None)
    return record


def split_rows(record: ImportToken, *, add_unknown_units: bool) -> tuple[list[ParsedRow], list[ParsedRow], list[str]]:
    """按 add_unknown_units 拆分：待落库行 / 跳过行 / 新增单位。"""
    if add_unknown_units:
        return list(record.rows), [], sorted({r.unit for r in record.rows if r.unknown_unit and r.unit})
    keep = [r for r in record.rows if not r.unknown_unit]
    skipped = [r for r in record.rows if r.unknown_unit]
    return keep, skipped, []
