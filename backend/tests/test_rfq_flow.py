"""端到端闭环用例：建单 → 发布 → 保存/改价/提交 → 进度 → 比价 → 汇总 → 导出。

金额全部按人工核算值断言（PRD TC-23 ~ TC-27、TC-36 ~ TC-41）。
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from .conftest import auth


def _lines_of(detail: dict) -> list[dict]:
    return detail["lines"]


def test_full_flow(client, buyer_token, supplier_token, supplier2_token, rfq_factory, accounts):
    # ---------- 1. 建单（3 行明细 / 2 家供货商） ----------
    detail = rfq_factory(lines=3, supplier_indexes=[0, 1], title="闭环询价单")
    rfq_id = detail["id"]
    assert detail["status"] == "draft"
    assert detail["rfq_no"].startswith("RFQ")
    assert len(detail["lines"]) == 3
    assert len(detail["assignees"]) == 2
    assert all(a["line_count"] == 3 for a in detail["assignees"])
    uids = [line["line_uid"] for line in detail["lines"]]

    # ---------- 2. 发布校验：截止时间为过去 -> 阻断并逐条列出 ----------
    bad = client.put(
        f"/api/rfqs/{rfq_id}",
        json={
            "title": "闭环询价单",
            "project": "测试项目",
            "deadline": "2020-01-01T00:00:00",
            "supplier_ids": [accounts["suppliers"][0]["id"]],
            "lines": [{"material_name": "物料1", "unit": "个", "quantity": "10"}],
        },
        headers=auth(buyer_token),
    )
    assert bad.status_code == 200, bad.text
    publish = client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token))
    assert publish.status_code == 400
    body = publish.json()
    assert body["code"] == 40001
    assert "报价截止时间：必须晚于当前时间" in body["message"]

    # 无明细 / 无供货商
    empty = client.put(
        f"/api/rfqs/{rfq_id}",
        json={"title": "闭环询价单", "project": "测试项目", "deadline": None, "supplier_ids": [], "lines": []},
        headers=auth(buyer_token),
    )
    assert empty.status_code == 200
    publish = client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token))
    assert publish.status_code == 400
    message = publish.json()["message"]
    assert "报价截止时间：必填" in message
    assert "明细：至少需要 1 行明细" in message
    assert "受邀供货商：请至少选择一家供货商" in message

    # ---------- 3. 恢复内容并发布 ----------
    from datetime import datetime, timedelta

    fix = client.put(
        f"/api/rfqs/{rfq_id}",
        json={
            "title": "闭环询价单",
            "project": "测试项目",
            "deadline": (datetime.now() + timedelta(days=5)).strftime("%Y-%m-%dT%H:%M:%S"),
            "supplier_ids": [accounts["suppliers"][0]["id"], accounts["suppliers"][1]["id"]],
            "lines": [
                {"material_name": f"物料{seq}", "spec": f"规格{seq}", "unit": "个", "quantity": str(seq * 10)}
                for seq in (1, 2, 3)
            ],
        },
        headers=auth(buyer_token),
    )
    assert fix.status_code == 200
    publish = client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token))
    assert publish.status_code == 200, publish.text
    published = publish.json()["data"]
    assert published["status"] == "published"
    assert len(published["assignees"]) == 2
    assert all(a["status"] == "pending" for a in published["assignees"])
    uids = [line["line_uid"] for line in published["lines"]]

    # 已发布后禁止改明细（40900）
    blocked = client.put(
        f"/api/rfqs/{rfq_id}",
        json={
            "title": "闭环询价单",
            "project": "测试项目",
            "deadline": (datetime.now() + timedelta(days=5)).strftime("%Y-%m-%dT%H:%M:%S"),
            "supplier_ids": [accounts["suppliers"][0]["id"]],
            "lines": [],
        },
        headers=auth(buyer_token),
    )
    assert blocked.status_code == 409
    assert blocked.json()["code"] == 40900

    # ---------- 4. 供货商 1：保存部分行 ----------
    sf = client.get(f"/api/supplier/rfqs/{rfq_id}", headers=auth(supplier_token))
    assert sf.status_code == 200
    page = sf.json()["data"]
    assert page["my_status"] == "pending"
    assert page["readonly"] is False
    assert all(line["quote"] is None for line in page["lines"])

    save = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={
            "submit": False,
            "lines": [
                {"line_uid": uids[0], "price_incl_tax": "113.00", "tax_included": True, "tax_rate": "0.13"}
            ],
        },
        headers=auth(supplier_token),
    )
    assert save.status_code == 200, save.text
    saved = save.json()["data"]
    assert saved["my_status"] == "partial"
    assert saved["quoted_line_count"] == 1
    assert saved["line_count"] == 3
    # 不含税单价 = 113 / 1.13 = 100.00（TC-27）
    assert saved["lines"][0]["quote"]["price_excl_tax"] == "100.00"

    history = client.get(
        f"/api/supplier/quotes/history?rfq_id={rfq_id}", headers=auth(supplier_token)
    ).json()["data"]
    assert history["meta"]["total"] == 1
    first = history["items"][0]
    assert first["version"] == 1 and first["action"] == "save"
    assert first["snapshot"]["price_incl_tax"] is None

    # ---------- 5. 改价：113 → 95，版本 2 且快照为改前值 ----------
    change = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={
            "submit": False,
            "lines": [
                {"line_uid": uids[0], "price_incl_tax": "95.00", "tax_included": True, "tax_rate": "0.13"}
            ],
        },
        headers=auth(supplier_token),
    )
    assert change.status_code == 200
    history = client.get(
        f"/api/supplier/quotes/history?rfq_id={rfq_id}&line_uid={uids[0]}",
        headers=auth(supplier_token),
    ).json()["data"]
    versions = sorted(item["version"] for item in history["items"])
    assert versions == [1, 2]
    latest = history["items"][0]
    assert latest["version"] == 2
    assert latest["snapshot"]["price_incl_tax"] == "113.00"
    assert "price_incl_tax" in latest["changed_fields"]

    # ---------- 6. 提交校验：仅 1/3 行已填 ----------
    fail = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={"submit": True, "lines": []},
        headers=auth(supplier_token),
    )
    assert fail.status_code == 400
    assert fail.json()["code"] == 40001
    assert "第2行 含税单价未填写" in fail.json()["message"]

    # ---------- 7. 提交（供1: 95/22/33） ----------
    submit = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={
            "submit": True,
            "lines": [
                {"line_uid": uids[0], "price_incl_tax": "95.00", "tax_included": True, "tax_rate": "0.13", "moq": "1", "lead_time_days": 4},
                {"line_uid": uids[1], "price_incl_tax": "22.00", "tax_included": True, "tax_rate": "0.13", "moq": "1", "lead_time_days": 4},
                {"line_uid": uids[2], "price_incl_tax": "33.00", "tax_included": True, "tax_rate": "0.13", "moq": "1", "lead_time_days": 4},
            ],
        },
        headers=auth(supplier_token),
    )
    assert submit.status_code == 200, submit.text
    assert submit.json()["data"]["my_status"] == "quoted"
    assert submit.json()["data"]["quoted_line_count"] == 3

    # 提交后仍可修改（未截止，TC-25）
    again = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={
            "submit": False,
            "lines": [{"line_uid": uids[2], "price_incl_tax": "33.00", "tax_included": True, "tax_rate": "0.13", "remark": "提交后修改验证"}],
        },
        headers=auth(supplier_token),
    )
    assert again.status_code == 200
    assert again.json()["data"]["my_status"] == "quoted"

    # ---------- 8. 供货商 2：提交（100/25/30） ----------
    # 用 PRD 7.3 的单列提交入口 POST /rfqs/{id}/submit（等价 PUT + submit=true）
    submit2 = client.post(
        f"/api/supplier/rfqs/{rfq_id}/submit",
        json={
            "lines": [
                {"line_uid": uids[0], "price_incl_tax": "100.00", "tax_included": True, "tax_rate": "0.13", "moq": "3", "lead_time_days": 6},
                {"line_uid": uids[1], "price_incl_tax": "25.00", "tax_included": True, "tax_rate": "0.13", "moq": "3", "lead_time_days": 6},
                {"line_uid": uids[2], "price_incl_tax": "30.00", "tax_included": True, "tax_rate": "0.13", "moq": "3", "lead_time_days": 6},
            ],
        },
        headers=auth(supplier2_token),
    )
    assert submit2.status_code == 200, submit2.text
    assert submit2.json()["data"]["my_status"] == "quoted"

    # ---------- 9. 我方进度 ----------
    progress = client.get(f"/api/rfqs/{rfq_id}/progress", headers=auth(buyer_token)).json()["data"]
    assert len(progress) == 2
    assert all(item["status"] == "quoted" for item in progress)
    assert all(item["quoted_line_count"] == 3 for item in progress)

    listing = client.get("/api/rfqs?page_size=100", headers=auth(buyer_token)).json()["data"]
    row = next(item for item in listing["items"] if item["id"] == rfq_id)
    assert row["line_count"] == 3
    assert row["supplier_count"] == 2
    assert row["quoted_count"] == 2

    # ---------- 10. 比价报表（视图 A，人工核算） ----------
    comparison = client.get(
        f"/api/reports/rfq-comparison?rfq_id={rfq_id}", headers=auth(buyer_token)
    ).json()["data"]
    assert comparison["status"] == "published"
    rows = {row["seq"]: row for row in comparison["rows"]}

    assert rows[1]["prices"]["1"] == "95.00" and rows[1]["prices"]["2"] == "100.00"
    assert rows[1]["best_price"] == "95.00"
    assert rows[1]["best_supplier_name"] == accounts["suppliers"][0]["name"]
    assert rows[1]["max_price"] == "100.00"
    assert rows[1]["avg_price"] == "97.50"
    assert rows[1]["diff_rate"] == "5.3%"  # (100-95)/95 = 5.26% → 一位小数
    assert rows[1]["quoted_count"] == 2 and rows[1]["required_count"] == 2

    assert rows[2]["best_price"] == "22.00"
    assert rows[2]["best_supplier_name"] == accounts["suppliers"][0]["name"]
    assert rows[2]["diff_rate"] == "13.6%"  # (25-22)/22 = 13.64%

    assert rows[3]["best_price"] == "30.00"
    assert rows[3]["best_supplier_name"] == accounts["suppliers"][1]["name"]
    assert rows[3]["diff_rate"] == "10.0%"  # (33-30)/30 = 10%

    # 合计：最优 95*10 + 22*20 + 30*30 = 2290；最高 100*10 + 25*20 + 33*30 = 2490
    assert comparison["total_best_amount"] == "2290.00"
    assert comparison["total_max_amount"] == "2490.00"
    assert comparison["saved_amount"] == "200.00"
    assert comparison["saved_rate"] == "8.0%"  # 200/2490 = 8.03%

    # 我方报价明细接口按 line_uid 聚合，与比价一致
    quotes = client.get(f"/api/rfqs/{rfq_id}/quotes", headers=auth(buyer_token)).json()["data"]
    assert quotes["total_best_amount"] == "2290.00"

    # ---------- 11. 汇总报表（视图 B） ----------
    summary = client.get(
        f"/api/reports/summary?rfq_id={rfq_id}", headers=auth(buyer_token)
    ).json()["data"]
    assert summary["rfq_count"] == 1
    assert summary["line_count"] == 3
    assert summary["invited_assignments"] == 2
    assert summary["quoted_assignments"] == 2
    assert summary["response_rate"] == "100.0%"
    assert summary["best_amount"] == "2290.00"
    by_supplier = {row["supplier_id"]: row for row in summary["by_supplier"]}
    first_supplier = by_supplier[accounts["suppliers"][0]["id"]]
    # 供1 总额 = 95*10 + 22*20 + 33*30 = 950 + 440 + 990 = 2380
    assert first_supplier["amount"] == "2380.00"
    assert first_supplier["response_rate"] == "100.0%"
    assert first_supplier["avg_lead_time"] == "4.0"
    assert first_supplier["avg_moq"] == "1.000"
    second_supplier = by_supplier[accounts["suppliers"][1]["id"]]
    # 供2 总额 = 100*10 + 25*20 + 30*30 = 1000 + 500 + 900 = 2400
    assert second_supplier["amount"] == "2400.00"
    assert summary["by_project"][0]["project"] == "测试项目"
    assert summary["by_project"][0]["line_count"] == 3
    assert summary["by_project"][0]["coverage_rate"] == "100.0%"

    # ---------- 12. 导出 xlsx ----------
    export = client.get(f"/api/rfqs/{rfq_id}/export", headers=auth(buyer_token))
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in export.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(export.content))
    sheet = workbook["询价单"]
    header_row = next(
        row for row in range(1, sheet.max_row + 1) if sheet.cell(row=row, column=1).value == "序号"
    )
    data_rows = [
        row
        for row in range(header_row + 1, sheet.max_row + 1)
        if isinstance(sheet.cell(row=row, column=1).value, int)
    ]
    assert len(data_rows) == 3
    # 最低价底色高亮：每行一个
    highlighted = [
        (row, col)
        for row in data_rows
        for col in range(7, 7 + 2 + 6)
        if sheet.cell(row=row, column=col).fill.fgColor.rgb == "00FFF2CC"
    ]
    assert len(highlighted) == 3
    # 价差率 / 节省率在导出中同样是百分数字符串
    diff_col = 6 + 2 + 5
    assert sheet.cell(row=data_rows[0], column=diff_col).value == "5.3%"
    assert "报价进度" in workbook.sheetnames
    workbook.close()

    comparison_export = client.get(
        f"/api/reports/rfq-comparison/export?rfq_id={rfq_id}", headers=auth(buyer_token)
    )
    assert comparison_export.status_code == 200
    workbook = load_workbook(io.BytesIO(comparison_export.content))
    assert workbook.sheetnames == ["比价明细", "供货商"]
    workbook.close()

    summary_export = client.get("/api/reports/summary/export", headers=auth(buyer_token))
    assert summary_export.status_code == 200
    workbook = load_workbook(io.BytesIO(summary_export.content))
    assert "按供货商" in workbook.sheetnames and "按项目" in workbook.sheetnames
    workbook.close()

    # ---------- 13. 截止后锁定（40301） ----------
    closed = client.post(f"/api/rfqs/{rfq_id}/close", headers=auth(buyer_token))
    assert closed.status_code == 200
    assert closed.json()["data"]["status"] == "closed"

    late = client.put(
        f"/api/supplier/rfqs/{rfq_id}/quotes",
        json={"submit": False, "lines": [{"line_uid": uids[0], "price_incl_tax": "1.00"}]},
        headers=auth(supplier_token),
    )
    assert late.status_code == 403
    assert late.json()["code"] == 40301

    readonly = client.get(f"/api/supplier/rfqs/{rfq_id}", headers=auth(supplier_token)).json()["data"]
    assert readonly["readonly"] is True

    # 已截止单据禁止删除（40900）
    assert client.delete(f"/api/rfqs/{rfq_id}", headers=auth(buyer_token)).json()["code"] == 40900


def test_draft_allows_incomplete_lines_but_publish_blocks(client, buyer_token, accounts):
    """PRD FR-RFQ-09 / TC-09 与 FR-RFQ-10。

    草稿允许明细不完整（只提示不阻断），完整性由**发布校验**逐条拦截。
    对应 lead 对 schemas.RfqLineIn（material_name/unit/quantity 放开为可选）与
    api/rfq.py::_replace_lines（空值落成空串/0）的修复，防止被回退。
    """
    from datetime import datetime, timedelta

    deadline = (datetime.now() + timedelta(days=3)).strftime("%Y-%m-%dT%H:%M:%S")
    created = client.post(
        "/api/rfqs",
        headers=auth(buyer_token),
        json={
            "title": "不完整明细草稿",
            "project": "测试项目",
            "deadline": deadline,
            "supplier_ids": [accounts["suppliers"][0]["id"]],
            "lines": [
                {"material_name": "完整行", "unit": "个", "quantity": "10"},
                {},  # 完全不完整：材料名称/单位/采购数量全空
                {"material_name": "", "unit": "", "quantity": None},
                {"material_name": "只填名称"},
            ],
        },
    )
    assert created.status_code == 200, created.text  # 草稿保存成功，不被阻断
    detail = created.json()["data"]
    rfq_id = detail["id"]
    assert len(detail["lines"]) == 4
    assert detail["lines"][1]["material_name"] == ""
    assert detail["lines"][1]["unit"] == ""
    assert detail["lines"][1]["quantity"] == "0.000"

    # 数量 0/负数仍被入参校验挡住（保留 gt=0 约束）
    bad_qty = client.post(
        "/api/rfqs",
        headers=auth(buyer_token),
        json={
            "title": "数量非法",
            "project": "测试项目",
            "deadline": deadline,
            "supplier_ids": [],
            "lines": [{"material_name": "物料", "unit": "个", "quantity": "0"}],
        },
    )
    assert bad_qty.status_code == 400
    assert bad_qty.json()["code"] == 40001

    # 发布被逐条拦截
    publish = client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token))
    assert publish.status_code == 400
    body = publish.json()
    assert body["code"] == 40001
    message = body["message"]
    assert "第2行 材料名称：必填" in message
    assert "第2行 单位：必填" in message
    assert "第2行 采购数量：必须大于 0" in message
    assert "第3行 材料名称：必填" in message
    assert "第4行 单位：必填" in message
    assert "第4行 采购数量：必须大于 0" in message

    # 补齐后即可发布
    fixed = client.put(
        f"/api/rfqs/{rfq_id}",
        headers=auth(buyer_token),
        json={
            "title": "不完整明细草稿",
            "project": "测试项目",
            "deadline": deadline,
            "supplier_ids": [accounts["suppliers"][0]["id"]],
            "lines": [
                {"material_name": f"物料{seq}", "unit": "个", "quantity": str(seq)}
                for seq in (1, 2, 3, 4)
            ],
        },
    )
    assert fixed.status_code == 200, fixed.text
    published = client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token))
    assert published.status_code == 200, published.text
    assert published.json()["data"]["status"] == "published"
