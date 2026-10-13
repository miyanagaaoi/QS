"""Excel 导入用例：预校验精确到行列、.xls 兼容、未知单位、上限、模式、模板（TC-12 ~ TC-15）。"""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from .conftest import DOC_SAMPLE_XLS, auth

HEADER = ["序号", "材料名称", "规格型号", "单位", "项目", "采购数量"]


def _xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价明细"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _preview(client, token, rfq_id: int, filename: str, content: bytes):
    return client.post(
        f"/api/rfqs/import/preview?rfq_id={rfq_id}",
        headers=auth(token),
        files={"file": (filename, content, "application/octet-stream")},
    )


def _confirm(client, token, payload: dict):
    return client.post("/api/rfqs/import/confirm", headers=auth(token), json=payload)


def test_errors_are_reported_by_row_and_column(client, buyer_token, rfq_factory):
    detail = rfq_factory(lines=1, title="导入错误校验")
    rfq_id = detail["id"]
    content = _xlsx(
        [
            [1, "正常行", "A4", "个", "测试项目", 10],
            [2, "数量非数字", "X", "个", "测试项目", "abc"],
            [3, "数量为零", "X", "个", "测试项目", 0],
            [4, "单位不存在", "X", "箱包", "测试项目", 5],
            [5, "超长" * 60, "X", "个", "测试项目", 5],
            [None, None, None, None, None, None],
            [6, "末尾正常行", "Y", "台", "测试项目", 7],
        ]
    )
    response = _preview(client, buyer_token, rfq_id, "errors.xlsx", content)
    assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["total_rows"] == 7
    assert result["valid_rows"] == 2
    assert result["error_rows"] == 5
    assert result["token"] is None  # 硬错误不签发 token，未落库

    problems = {(item["row"], item["column"]) for item in result["errors"]}
    assert (3, "采购数量") in problems
    assert (4, "采购数量") in problems
    assert (5, "单位") in problems
    assert (6, "材料名称") in problems
    assert (7, None) in problems
    reasons = " ".join(item["reason"] for item in result["errors"])
    assert "必须为大于 0 的数字" in reasons
    assert "文本超长" in reasons
    assert "整行内容为空" in reasons

    # 未落库：明细行数不变
    after = client.get(f"/api/rfqs/{rfq_id}", headers=auth(buyer_token)).json()["data"]
    assert len(after["lines"]) == 1


def test_xls_ole2_compatibility_and_confirm(client, buyer_token, rfq_factory):
    """doc/示例.xls 是 OLE2 容器（伪装 .xls），必须能直接读取并落库（TC-14）。"""
    assert DOC_SAMPLE_XLS.exists(), DOC_SAMPLE_XLS
    detail = rfq_factory(lines=1, title="xls 兼容")
    rfq_id = detail["id"]

    response = _preview(client, buyer_token, rfq_id, "示例.xls", DOC_SAMPLE_XLS.read_bytes())
    assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["total_rows"] == 1
    assert result["valid_rows"] == 1
    assert result["errors"] == []
    assert result["token"]

    confirmed = _confirm(
        client, buyer_token, {"token": result["token"], "mode": "append", "add_unknown_units": False}
    )
    assert confirmed.status_code == 200, confirmed.text
    data = confirmed.json()["data"]
    assert data["imported_rows"] == 1
    assert data["total_lines"] == 2

    # token 一次性
    reused = _confirm(
        client, buyer_token, {"token": result["token"], "mode": "append", "add_unknown_units": False}
    )
    assert reused.status_code == 400
    assert reused.json()["code"] == 40001


def test_unknown_unit_skip_and_add(client, buyer_token, rfq_factory):
    detail = rfq_factory(lines=1, title="未知单位")
    rfq_id = detail["id"]
    unit = "测打A"
    content = _xlsx([[1, "未知单位物料", "S", unit, "测试项目", 9]])

    response = _preview(client, buyer_token, rfq_id, "unknown.xlsx", content)
    result = response.json()["data"]
    assert result["unknown_units"] == [unit]
    assert result["token"]  # 未知单位属于软错误，仍签发 token

    skipped = _confirm(
        client, buyer_token, {"token": result["token"], "mode": "append", "add_unknown_units": False}
    )
    assert skipped.status_code == 200, skipped.text
    assert skipped.json()["data"]["imported_rows"] == 0
    assert skipped.json()["data"]["skipped_rows"] == 1

    response = _preview(client, buyer_token, rfq_id, "unknown.xlsx", content)
    added = _confirm(
        client, buyer_token,
        {"token": response.json()["data"]["token"], "mode": "append", "add_unknown_units": True},
    )
    assert added.status_code == 200, added.text
    assert added.json()["data"]["new_units"] == [unit]
    assert added.json()["data"]["imported_rows"] == 1

    # 第二次导入时该单位已在字典内，可直接成功
    response = _preview(client, buyer_token, rfq_id, "unknown.xlsx", content)
    assert response.json()["data"]["unknown_units"] == []
    assert response.json()["data"]["valid_rows"] == 1


def test_replace_mode_and_row_limit(client, buyer_token, rfq_factory):
    detail = rfq_factory(lines=3, title="替换模式")
    rfq_id = detail["id"]
    response = _preview(
        client, buyer_token, rfq_id, "replace.xlsx", _xlsx([[1, "替换A", "", "个", "测试项目", 3]])
    )
    confirmed = _confirm(
        client, buyer_token,
        {"token": response.json()["data"]["token"], "mode": "replace", "add_unknown_units": False},
    )
    assert confirmed.json()["data"]["total_lines"] == 1
    lines = client.get(f"/api/rfqs/{rfq_id}", headers=auth(buyer_token)).json()["data"]["lines"]
    assert [line["material_name"] for line in lines] == ["替换A"]

    # 单次上限 2000 行（TC-15）
    too_many = _xlsx([[index, f"物料{index}", "", "个", "测试项目", 1] for index in range(1, 2002)])
    response = _preview(client, buyer_token, rfq_id, "too_many.xlsx", too_many)
    assert response.status_code == 400
    assert "超过单次导入上限" in response.json()["message"]


def test_template_and_bad_extension(client, buyer_token, rfq_factory):
    response = client.get("/api/rfqs/import/template", headers=auth(buyer_token))
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["询价明细", "填写说明"]
    assert [cell.value for cell in workbook["询价明细"][1]] == HEADER
    workbook.close()

    detail = rfq_factory(lines=1, title="非法扩展名")
    bad = _preview(client, buyer_token, detail["id"], "bad.txt", b"hello")
    assert bad.status_code == 400
    assert bad.json()["code"] == 40001


def test_import_blocked_after_publish(client, buyer_token, rfq_factory):
    detail = rfq_factory(lines=1, title="发布后导入")
    rfq_id = detail["id"]
    assert client.post(f"/api/rfqs/{rfq_id}/publish", headers=auth(buyer_token)).json()["code"] == 0
    response = _preview(
        client, buyer_token, rfq_id, "any.xlsx", _xlsx([[1, "物料", "", "个", "测试项目", 1]])
    )
    assert response.status_code == 409
    assert response.json()["code"] == 40900
