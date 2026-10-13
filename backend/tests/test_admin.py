"""系统管理模块（task-5 / task-7）测试。

覆盖：最后一个管理员保护、批量导入校验与落库、默认税率唯一性、账号与日志不可删除、
重置密码后旧 token 立即失效、供货商账号清单导出（不含密码）、导入错误报告导出。

隔离策略：
- 每个用例独立创建临时 SQLite 库（tmp_path），通过 dependency_overrides 覆盖 get_db，
  **完全不触碰共享库 backend/data/xunjia.db**；
- 不使用 TestClient 的上下文管理器，因此不会触发 lifespan（不建表、不写种子到生产库）；
- XUNJIA_BCRYPT_ROUNDS=4 降低哈希轮数，加快用例执行。

运行：cd backend && ./.venv/Scripts/python.exe -m pytest tests/test_admin.py -q
"""
from __future__ import annotations

import io
import os
import pathlib
import sys
from datetime import datetime
from decimal import Decimal
from urllib.parse import quote

import pytest

os.environ.setdefault("XUNJIA_BCRYPT_ROUNDS", "4")

BACKEND_DIR = pathlib.Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:  # 保证 `import app` 可用（不依赖 pytest rootdir 行为）
    sys.path.insert(0, str(BACKEND_DIR))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import create_engine, func, select  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app.database import get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models import AuditLog, Base, Supplier, TaxDict, UnitDict, User  # noqa: E402
from app.security import hash_password  # noqa: E402

ADMIN_PWD = "Admin@2026"
BUYER_PWD = "Buyer@2026"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_HEADERS = ["供货商名称", "联系人", "联系电话", "邮箱", "地址", "登录账号", "状态", "最后登录时间", "创建时间"]
ERROR_REPORT_HEADERS = ["Excel 行号", "供货商名称", "登录账号", "失败原因"]


class Ctx:
    def __init__(self, client, session_factory, admin_id, db_path):
        self.client = client
        self.session = session_factory
        self.admin_id = admin_id
        self.db_path = db_path

    def query(self, stmt):
        with self.session() as s:
            return s.execute(stmt).all()

    def scalar(self, stmt):
        with self.session() as s:
            return s.scalar(stmt)


@pytest.fixture()
def ctx(tmp_path):
    db_path = tmp_path / "test_admin.db"
    engine = create_engine(
        f"sqlite:///{db_path.as_posix()}", connect_args={"check_same_thread": False}, future=True
    )
    Base.metadata.create_all(bind=engine)
    SessionFactory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    with SessionFactory() as s:
        admin = User(
            username="admin",
            password_hash=hash_password(ADMIN_PWD),
            real_name="系统管理员",
            role="admin",
            status=1,
            must_change_password=False,
        )
        buyer = User(
            username="buyer01",
            password_hash=hash_password(BUYER_PWD),
            real_name="采购员甲",
            role="buyer",
            status=1,
            must_change_password=False,
        )
        s.add_all([admin, buyer])
        for i, name in enumerate(["个", "台", "包"], start=1):
            s.add(UnitDict(name=name, sort=i, status=1))
        s.add(TaxDict(rate=Decimal("0.1300"), label="13%", is_default=True, status=1, sort=1))
        s.add(TaxDict(rate=Decimal("0.0900"), label="9%", is_default=False, status=1, sort=2))
        s.commit()
        admin_id = admin.id

    def override_get_db():
        db = SessionFactory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)  # 不用 with：避免触发 lifespan 写入生产库
    try:
        yield Ctx(client, SessionFactory, admin_id, db_path)
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def _hdr(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _login(client, username: str, password: str) -> str:
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    assert r.json()["code"] == 0
    return r.json()["data"]["token"]


def _admin_token(ctx) -> str:
    return _login(ctx.client, "admin", ADMIN_PWD)


def _error_rows(resp_json) -> list[tuple[int, str | None, str]]:
    return [(e["row"], e["column"], e["reason"]) for e in resp_json["data"]["errors"]]


def _build_import_xlsx(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "供货商账号"
    ws.append(["供货商名称", "联系人", "联系电话", "邮箱", "地址", "登录账号", "初始密码"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_rows(content: bytes) -> list[list]:
    """用 openpyxl 复读接口返回的 xlsx，返回所有行（含表头）。"""
    wb = load_workbook(io.BytesIO(content))
    try:
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _xlsx_flat_text(content: bytes) -> str:
    return " | ".join(str(c) for row in _xlsx_rows(content) for c in row if c is not None)


# --------------------------------------------------------------------------
# 1. FR-ACC-06 最后一个启用管理员保护
# --------------------------------------------------------------------------
def test_last_enabled_admin_cannot_be_disabled_or_demoted(ctx):
    token = _admin_token(ctx)

    r = ctx.client.post(f"/api/admin/users/{ctx.admin_id}/toggle-status", headers=_hdr(token))
    assert r.status_code == 409, r.text
    body = r.json()
    assert body["code"] == 40900
    assert "最后一个启用的管理员" in body["message"]
    assert body["data"] is None

    r = ctx.client.put(
        f"/api/admin/users/{ctx.admin_id}", json={"role": "buyer"}, headers=_hdr(token)
    )
    assert r.status_code == 409, r.text
    assert r.json()["code"] == 40900
    assert "最后一个启用的管理员" in r.json()["message"]

    # 拒绝后原账号必须保持启用且仍为管理员
    user = ctx.scalar(select(User).where(User.id == ctx.admin_id))
    assert (user.status, user.role) == (1, "admin")

    # 拒绝动作必须留痕（FR-LOG-01）
    rejected = ctx.query(
        select(AuditLog.action, AuditLog.summary).where(AuditLog.summary.like("%最后一个启用的管理员%"))
    )
    assert rejected, "拒绝停用最后一个管理员时必须写入操作日志"
    assert any(a == "user.toggle_status" for a, _ in rejected)

    # 新增第二个启用管理员后，原管理员可以正常停用（说明规则是"最后一个"而非一律禁止）
    r = ctx.client.post(
        "/api/admin/users",
        json={
            "username": "admin02",
            "real_name": "第二管理员",
            "role": "admin",
            "initial_password": "Admin@2027",
        },
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    assert r.json()["data"]["must_change_password"] is False  # 需求变更：新建默认不强制改密（新建我方账号）

    r = ctx.client.post(f"/api/admin/users/{ctx.admin_id}/toggle-status", headers=_hdr(token))
    assert r.status_code == 200, r.text
    assert r.json()["data"]["status"] == 0


# --------------------------------------------------------------------------
# 2. FR-ACC-10~13 批量导入：预校验行级错误 + 确认落库
# --------------------------------------------------------------------------
def test_supplier_import_preview_reports_row_errors_and_confirm_persists_only_valid(ctx):
    token = _admin_token(ctx)
    # 预先存在的账号与档案，用于触发"系统内已存在""档案已绑定"两类错误
    with ctx.session() as s:
        bound_supplier = Supplier(name="既有供货商", contact_name="老张", status=1)
        free_supplier = Supplier(name="复用档案", contact_name="老王", status=1)
        s.add_all([bound_supplier, free_supplier])
        s.flush()
        sp_id = free_supplier.id
        s.add(
            User(
                username="exists01",
                password_hash=hash_password("Exists@2026"),
                role="supplier",
                supplier_id=bound_supplier.id,
                status=1,
                must_change_password=False,
            )
        )
        s.commit()

    rows = [
        ["导入A", "张三", "13800000001", "a@qa.example", "", "exists01", "Abc12345"],  # 账号已存在
        ["导入B", "", "13800000002", "b@qa.example", "", "imp_ok0", "Abc12345"],  # 必填缺失
        ["导入C", "李四", "13800000003", "not-an-email", "", "imp_badmail", "Abc12345"],  # 邮箱格式
        ["导入D", "王五", "13800000004", "d@qa.example", "", "imp_ok1", "Abc12345"],  # 合法
        ["导入E", "赵六", "13800000005", "e@qa.example", "", "imp_ok1", "Abc12345"],  # 账号文件内重复
        ["导入D", "钱七", "13800000006", "f@qa.example", "", "imp_ok2", "Abc12345"],  # 名称文件内重复
        ["导入F", "孙八", "12345", "g@qa.example", "", "imp_ok3", "Abc12345"],  # 电话格式
        ["导入G", "周九", "13800000009", "h@qa.example", "", "imp_ok4", "abcdefgh"],  # 密码强度
        ["复用档案", "吴十", "13800000010", "i@qa.example", "", "imp_ok5", "Abc12345"],  # 合法（复用）
        ["既有供货商", "郑十一", "13800000011", "j@qa.example", "", "imp_ok6", "Abc12345"],  # 档案已绑定
    ]
    raw = _build_import_xlsx(rows)

    users_before = ctx.scalar(select(func.count()).select_from(User))
    suppliers_before = ctx.scalar(select(func.count()).select_from(Supplier))

    r = ctx.client.post(
        "/api/admin/supplier-accounts/import",
        files={
            "file": (
                "err.xlsx",
                raw,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert data["total"] == 10
    assert data["valid"] == 2
    assert data["invalid"] == 8
    assert data["token"]

    errors = _error_rows(r.json())
    by_row = {row: (column, reason) for row, column, reason in errors}
    assert by_row[2][1].startswith("登录账号已存在于系统中")
    assert by_row[3][0] == "联系人" and "必填" in by_row[3][1]
    assert by_row[4][0] == "邮箱" and "格式不正确" in by_row[4][1]
    assert "文件内重复" in by_row[6][1]
    assert by_row[7][0] == "供货商名称" and "文件内重复" in by_row[7][1]
    assert by_row[8][0] == "联系电话"
    assert by_row[9][0] == "初始密码" and "字母" in by_row[9][1]
    assert by_row[11][1].startswith("该供货商档案已绑定账号")

    # 预校验阶段绝不落库
    assert ctx.scalar(select(func.count()).select_from(User)) == users_before
    assert ctx.scalar(select(func.count()).select_from(Supplier)) == suppliers_before

    # 确认导入：仅 2 条合法行落库，其中一行复用已有档案
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import/confirm",
        json={"token": data["token"]},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    result = r.json()["data"]
    assert (result["success"], result["failed"]) == (2, 0)
    assert result["errors"] == []

    created = {u.username: u for (u,) in ctx.query(select(User).where(User.username.in_(["imp_ok1", "imp_ok5"])))}
    assert set(created) == {"imp_ok1", "imp_ok5"}
    for u in created.values():
        assert u.role == "supplier"
        assert u.must_change_password is False  # 需求变更：批量导入默认不强制改密
        assert u.status == 1
    assert created["imp_ok5"].supplier_id == sp_id  # 复用已有档案，不重复建档
    assert ctx.scalar(select(func.count()).select_from(Supplier)) == suppliers_before + 1

    # token 一次性，禁止重复落库
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import/confirm",
        json={"token": data["token"]},
        headers=_hdr(token),
    )
    assert r.status_code == 400
    assert r.json()["code"] == 40001


def test_supplier_import_rejects_non_xlsx_and_header_mismatch(ctx):
    token = _admin_token(ctx)
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import",
        files={"file": ("accounts.csv", b"a,b,c", "text/csv")},
        headers=_hdr(token),
    )
    assert r.status_code == 400 and r.json()["code"] == 40001

    wb = Workbook()
    wb.active.append(["甲", "乙"])
    buf = io.BytesIO()
    wb.save(buf)
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import",
        files={
            "file": (
                "bad.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=_hdr(token),
    )
    assert r.status_code == 400
    assert "表头缺少列" in r.json()["message"]


# --------------------------------------------------------------------------
# 3. FR-MST-02 默认税率唯一性
# --------------------------------------------------------------------------
def test_tax_default_is_unique_and_switchable(ctx):
    token = _admin_token(ctx)

    r = ctx.client.post(
        "/api/admin/taxes",
        json={"rate": "0.0500", "label": "5%", "is_default": True, "sort": 9},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    new_id = r.json()["data"]["id"]
    assert r.json()["data"]["is_default"] is True

    defaults = ctx.query(select(TaxDict.id).where(TaxDict.is_default.is_(True)))
    assert [row[0] for row in defaults] == [new_id], "设置新默认税率时必须清除旧默认项"
    assert ctx.scalar(select(func.count()).select_from(TaxDict).where(TaxDict.is_default.is_(True))) == 1

    # 默认税率不可停用，也不可直接取消默认
    r = ctx.client.put(f"/api/admin/taxes/{new_id}", json={"status": 0}, headers=_hdr(token))
    assert r.status_code == 409 and r.json()["code"] == 40900
    r = ctx.client.put(f"/api/admin/taxes/{new_id}", json={"is_default": False}, headers=_hdr(token))
    assert r.status_code == 409 and r.json()["code"] == 40900

    # 切回原 13%：必须仍然唯一
    old_id = ctx.scalar(select(TaxDict.id).where(TaxDict.rate == Decimal("0.1300")))
    r = ctx.client.put(f"/api/admin/taxes/{old_id}", json={"is_default": True}, headers=_hdr(token))
    assert r.status_code == 200, r.text
    assert [row[0] for row in ctx.query(select(TaxDict.id).where(TaxDict.is_default.is_(True)))] == [old_id]

    # 相同税率不允许重复新增
    r = ctx.client.post(
        "/api/admin/taxes", json={"rate": "0.0500", "label": "5% 重复"}, headers=_hdr(token)
    )
    assert r.status_code == 409 and r.json()["code"] == 40900

    # /api/meta/taxes 只返回启用项，且带 is_default
    tax = ctx.scalar(select(TaxDict).where(TaxDict.id == new_id))
    assert tax.status == 1
    r = ctx.client.get("/api/meta/taxes", headers=_hdr(token))
    assert r.status_code == 200
    assert sum(1 for t in r.json()["data"] if t["is_default"]) == 1


# --------------------------------------------------------------------------
# 4. FR-ACC-07 / FR-LOG-04 账号与日志不可删除
# --------------------------------------------------------------------------
def test_accounts_and_audit_logs_are_not_deletable(ctx):
    token = _admin_token(ctx)
    assert ctx.client.delete(f"/api/admin/users/{ctx.admin_id}", headers=_hdr(token)).status_code == 405
    assert ctx.client.delete("/api/admin/suppliers/1", headers=_hdr(token)).status_code == 405
    assert ctx.client.delete("/api/admin/units/1", headers=_hdr(token)).status_code == 405
    assert ctx.client.delete("/api/admin/taxes/1", headers=_hdr(token)).status_code == 405
    # 日志：只有 GET 路由，任何删除请求都不会命中（404），即系统不提供删除能力
    assert ctx.client.delete("/api/admin/audit-logs/1", headers=_hdr(token)).status_code in (404, 405)


# --------------------------------------------------------------------------
# 5. FR-ACC-05 重置密码后强制改密（旧会话在业务接口上失效）
# --------------------------------------------------------------------------
def test_reset_password_forces_change_and_blocks_business_api(ctx):
    token = _admin_token(ctx)
    r = ctx.client.post(
        "/api/admin/users",
        json={
            "username": "buyer02",
            "real_name": "采购员乙",
            "role": "buyer",
            "initial_password": "Init@2026",
        },
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    uid = r.json()["data"]["id"]
    assert r.json()["data"]["must_change_password"] is False  # 需求变更：新建默认不强制改密（新建我方账号）

    r = ctx.client.post(
        f"/api/admin/users/{uid}/reset-password",
        json={"new_password": "Reset@2026"},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    assert r.json()["data"]["must_change_password"] is True

    # 新密码可登录，但未改密前业务接口一律 40101
    new_token = _login(ctx.client, "buyer02", "Reset@2026")
    r = ctx.client.get("/api/meta/units", headers=_hdr(new_token))
    assert r.status_code == 401
    assert r.json()["code"] == 40101
    # 改密放行
    r = ctx.client.post(
        "/api/auth/change-password",
        json={"old_password": "Reset@2026", "new_password": "Changed@2026"},
        headers=_hdr(new_token),
    )
    assert r.status_code == 200, r.text
    # 改密前的旧 token 立即失效（JWT 密码指纹），必须重新登录
    r = ctx.client.get("/api/meta/units", headers=_hdr(new_token))
    assert r.status_code == 401, r.text
    assert r.json()["code"] == 40100
    assert "密码已变更" in r.json()["message"]

    fresh_token = _login(ctx.client, "buyer02", "Changed@2026")
    r = ctx.client.get("/api/meta/units", headers=_hdr(fresh_token))
    assert r.status_code == 200, r.text


def test_reset_password_invalidates_token_issued_before_reset(ctx):
    """D 项结案：管理员重置密码后，重置前签发的 token 访问业务接口必须 401/40100。"""
    admin_token = _admin_token(ctx)
    buyer_id = ctx.scalar(select(User.id).where(User.username == "buyer01"))

    before_token = _login(ctx.client, "buyer01", BUYER_PWD)
    assert ctx.client.get("/api/meta/units", headers=_hdr(before_token)).status_code == 200

    r = ctx.client.post(
        f"/api/admin/users/{buyer_id}/reset-password",
        json={"new_password": "Reset3@2026"},
        headers=_hdr(admin_token),
    )
    assert r.status_code == 200, r.text
    assert r.json()["data"]["must_change_password"] is True

    # 重置前签发的 token：立即失效（密码指纹不匹配）
    r = ctx.client.get("/api/meta/units", headers=_hdr(before_token))
    assert r.status_code == 401, r.text
    assert r.json()["code"] == 40100
    assert "密码已变更" in r.json()["message"]

    # 新密码可登录，但必须走强制改密流程
    after_token = _login(ctx.client, "buyer01", "Reset3@2026")
    r = ctx.client.get("/api/meta/units", headers=_hdr(after_token))
    assert r.status_code == 401
    assert r.json()["code"] == 40101


# --------------------------------------------------------------------------
# 8. FR-ACC-14 供货商账号清单导出（安全红线：不含密码/哈希）
# --------------------------------------------------------------------------
def _seed_export_accounts(ctx) -> None:
    with ctx.session() as s:
        sup_a = Supplier(
            name="导出档案A",
            contact_name="张三",
            contact_phone="13800000001",
            contact_email="a@export.example",
            address="导出路1号",
            status=1,
        )
        sup_b = Supplier(
            name="导出档案B",
            contact_name="李四",
            contact_phone="13800000002",
            contact_email="b@export.example",
            address="导出路2号",
            status=1,
        )
        s.add_all([sup_a, sup_b])
        s.flush()
        s.add(
            User(
                username="exp_alpha",
                password_hash=hash_password("SecPass@2026"),
                real_name="张三",
                role="supplier",
                supplier_id=sup_a.id,
                status=1,
                must_change_password=False,
                last_login_at=datetime(2026, 10, 13, 9, 30, 0),
            )
        )
        s.add(
            User(
                username="exp_beta",
                password_hash=hash_password("OtherPass@2026"),
                real_name="李四",
                role="supplier",
                supplier_id=sup_b.id,
                status=0,
                must_change_password=False,
                last_login_at=None,
            )
        )
        s.commit()


def test_supplier_account_export_xlsx_contains_fields_and_never_passwords(ctx):
    token = _admin_token(ctx)
    _seed_export_accounts(ctx)

    r = ctx.client.get("/api/admin/supplier-accounts/export", headers=_hdr(token))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX_MIME)
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert quote("供货商账号清单") in disposition  # RFC 5987 中文名
    assert disposition.endswith(".xlsx")

    rows = _xlsx_rows(r.content)
    assert rows[0] == EXPORT_HEADERS
    assert len(rows) == 3  # 表头 + 2 个供货商账号
    exported = {row[5]: row for row in rows[1:]}
    assert set(exported) == {"exp_alpha", "exp_beta"}
    assert exported["exp_alpha"][0] == "导出档案A"
    assert exported["exp_alpha"][2] == "13800000001"
    assert exported["exp_alpha"][5] == "exp_alpha"
    assert exported["exp_alpha"][6] == "启用"
    assert exported["exp_alpha"][7] == "2026-10-13 09:30:00"
    assert exported["exp_beta"][6] == "停用"
    assert exported["exp_beta"][7] in (None, "")  # 从未登录（空单元格 openpyxl 读回 None）

    # 安全红线：导出内容不得含明文密码、哈希片段或 password 字样
    flat = _xlsx_flat_text(r.content)
    assert "SecPass@2026" not in flat
    assert "OtherPass@2026" not in flat
    assert "$2b$" not in flat and "$2a$" not in flat
    assert "password" not in flat.lower()
    assert "hash" not in flat.lower()

    # 筛选参数与 GET /supplier-accounts 一致
    r = ctx.client.get(
        "/api/admin/supplier-accounts/export?keyword=exp_alpha&status=1", headers=_hdr(token)
    )
    rows = _xlsx_rows(r.content)
    assert len(rows) == 2 and rows[1][5] == "exp_alpha"

    r = ctx.client.get(
        "/api/admin/supplier-accounts/export?login_state=never&status=0", headers=_hdr(token)
    )
    rows = _xlsx_rows(r.content)
    assert len(rows) == 2 and rows[1][5] == "exp_beta"

    # 导出必须留痕
    actions = [a for (a,) in ctx.query(select(AuditLog.action).where(AuditLog.action == "supplier_account.export"))]
    assert len(actions) == 3

    # 无数据时只返回表头
    r = ctx.client.get("/api/admin/supplier-accounts/export?keyword=不存在的档案", headers=_hdr(token))
    assert len(_xlsx_rows(r.content)) == 1

    # 非管理员不可导出
    buyer_token = _login(ctx.client, "buyer01", BUYER_PWD)
    r = ctx.client.get("/api/admin/supplier-accounts/export", headers=_hdr(buyer_token))
    assert r.status_code == 403 and r.json()["code"] == 40300


# --------------------------------------------------------------------------
# 9. FR-ACC-12 导入失败明细导出
# --------------------------------------------------------------------------
def test_import_error_report_xlsx_matches_preview_errors(ctx):
    token = _admin_token(ctx)
    with ctx.session() as s:
        s.add(Supplier(name="错误报告既有档案", contact_name="老王", status=1))
        s.add(
            User(
                username="err_exists",
                password_hash=hash_password("Exists@2026"),
                role="supplier",
                status=1,
                must_change_password=False,
            )
        )
        s.commit()

    rows = [
        ["报告A", "张三", "13800000001", "a@err.example", "", "err_exists", "Abc12345"],  # 账号已存在
        ["报告B", "", "13800000002", "b@err.example", "", "rep_ok0", "Abc12345"],  # 必填缺失
        ["报告C", "李四", "13800000003", "bad-email", "", "rep_bad", "Abc12345"],  # 邮箱格式
        ["报告D", "王五", "13800000004", "d@err.example", "", "rep_ok1", "Abc12345"],  # 合法
        ["报告E", "赵六", "13800000005", "e@err.example", "", "rep_ok1", "Abc12345"],  # 账号文件内重复
    ]
    raw = _build_import_xlsx(rows)

    users_before = ctx.scalar(select(func.count()).select_from(User))
    suppliers_before = ctx.scalar(select(func.count()).select_from(Supplier))

    # 先用 /import 取同一份文件的错误清单，作为错误报告的口径基准
    preview = ctx.client.post(
        "/api/admin/supplier-accounts/import",
        files={"file": ("err.xlsx", raw, XLSX_MIME)},
        headers=_hdr(token),
    )
    assert preview.status_code == 200, preview.text
    preview_errors = _error_rows(preview.json())
    assert len(preview_errors) == 4

    r = ctx.client.post(
        "/api/admin/supplier-accounts/import/error-report",
        files={"file": ("err.xlsx", raw, XLSX_MIME)},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX_MIME)
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert quote("供货商账号导入错误报告") in disposition
    assert disposition.endswith(".xlsx")

    report = _xlsx_rows(r.content)
    assert report[0] == ERROR_REPORT_HEADERS
    assert len(report) == len(preview_errors) + 1  # 表头 + 与 /import 完全同数的错误行

    # 逐条对应：Excel 行号、供货商名称、登录账号、失败原因（含列名）
    by_row = {int(row[0]): row for row in report[1:]}
    assert set(by_row) == {2, 3, 4, 6}
    assert by_row[2][1] == "报告A" and by_row[2][2] == "err_exists"
    assert "[登录账号]" in by_row[2][3] and "已存在于系统中" in by_row[2][3]
    assert by_row[3][1] == "报告B" and "[联系人]" in by_row[3][3] and "必填" in by_row[3][3]
    assert by_row[4][1] == "报告C" and "[邮箱]" in by_row[4][3] and "格式不正确" in by_row[4][3]
    assert by_row[6][2] == "rep_ok1" and "[登录账号]" in by_row[6][3] and "文件内重复" in by_row[6][3]

    # 只校验不落库、不签发 token
    assert ctx.scalar(select(func.count()).select_from(User)) == users_before
    assert ctx.scalar(select(func.count()).select_from(Supplier)) == suppliers_before
    assert "token" not in r.text

    # 全部行合法 → 只有表头，且 audit summary 说明"无错误行"
    ok_raw = _build_import_xlsx(
        [["合法档案甲", "张三", "13800000001", "a@ok.example", "", "ok_user1", "Abc12345"]]
    )
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import/error-report",
        files={"file": ("ok.xlsx", ok_raw, XLSX_MIME)},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    report = _xlsx_rows(r.content)
    assert report == [ERROR_REPORT_HEADERS]
    assert ctx.scalar(select(func.count()).select_from(User)) == users_before

    latest = ctx.scalar(
        select(AuditLog.summary)
        .where(AuditLog.action == "supplier_account.import_error_report")
        .order_by(AuditLog.id.desc())
    )
    assert "无错误行" in latest

    # 非 xlsx / 表头缺失沿用同一套错误码
    r = ctx.client.post(
        "/api/admin/supplier-accounts/import/error-report",
        files={"file": ("a.csv", b"x,y", "text/csv")},
        headers=_hdr(token),
    )
    assert r.status_code == 400 and r.json()["code"] == 40001


# --------------------------------------------------------------------------
# 6. 权限与列表分页
# --------------------------------------------------------------------------
def test_admin_apis_require_admin_and_pagination_works(ctx):
    # 未登录 401
    assert ctx.client.get("/api/admin/users").status_code == 401

    buyer_token = _login(ctx.client, "buyer01", BUYER_PWD)
    r = ctx.client.get("/api/admin/users", headers=_hdr(buyer_token))
    assert r.status_code == 403 and r.json()["code"] == 40300
    # 我方人员可用 meta 下拉
    assert ctx.client.get("/api/meta/units", headers=_hdr(buyer_token)).status_code == 200
    assert ctx.client.get("/api/meta/taxes", headers=_hdr(buyer_token)).status_code == 200
    assert ctx.client.get("/api/meta/suppliers", headers=_hdr(buyer_token)).status_code == 200

    token = _admin_token(ctx)
    r = ctx.client.get("/api/admin/users?page=1&page_size=1", headers=_hdr(token))
    assert r.status_code == 200
    body = r.json()["data"]
    assert body["meta"] == {"page": 1, "page_size": 1, "total": 2}
    assert len(body["items"]) == 1
    assert body["items"][0]["username"] == "admin"

    r = ctx.client.get("/api/admin/users?keyword=采购&role=buyer", headers=_hdr(token))
    assert [i["username"] for i in r.json()["data"]["items"]] == ["buyer01"]

    # 账号不可删除 & 用户名唯一（FR-ACC-02）
    r = ctx.client.post(
        "/api/admin/users",
        json={
            "username": "buyer02",
            "real_name": "x",
            "role": "buyer",
            "initial_password": "Init@2026",
        },
        headers=_hdr(token),
    )
    assert r.status_code == 200
    r = ctx.client.post(
        "/api/admin/users",
        json={
            "username": "buyer02",
            "real_name": "x",
            "role": "buyer",
            "initial_password": "Init@2026",
        },
        headers=_hdr(token),
    )
    assert r.status_code == 409 and r.json()["code"] == 40900


# --------------------------------------------------------------------------
# 7. FR-ACC-09 供货商账号与档案一对一
# --------------------------------------------------------------------------
def test_supplier_account_binds_archive_one_to_one(ctx):
    token = _admin_token(ctx)
    r = ctx.client.post(
        "/api/admin/suppliers",
        json={"name": "新档案甲", "contact_name": "甲", "contact_phone": "13800000001"},
        headers=_hdr(token),
    )
    assert r.status_code == 200, r.text
    sid = r.json()["data"]["id"]

    payload = {
        "supplier_id": sid,
        "username": "sup_alpha",
        "initial_password": "Sup@2026x",
    }
    r = ctx.client.post("/api/admin/supplier-accounts", json=payload, headers=_hdr(token))
    assert r.status_code == 200, r.text
    assert r.json()["data"]["supplier_id"] == sid
    assert r.json()["data"]["must_change_password"] is False  # 需求变更：新建默认不强制改密（新建供货商账号）

    # 同一档案再绑一个账号 -> 40900
    r = ctx.client.post(
        "/api/admin/supplier-accounts",
        json={"supplier_id": sid, "username": "sup_beta", "initial_password": "Sup@2026x"},
        headers=_hdr(token),
    )
    assert r.status_code == 409 and r.json()["code"] == 40900

    # 不存在档案 -> 40400
    r = ctx.client.post(
        "/api/admin/supplier-accounts",
        json={"supplier_id": 99999, "username": "sup_gamma", "initial_password": "Sup@2026x"},
        headers=_hdr(token),
    )
    assert r.status_code == 404 and r.json()["code"] == 40400

    # 停用档案 -> 同步停用绑定账号（FR-MST-04）并且账号无法登录
    r = ctx.client.post(f"/api/admin/suppliers/{sid}/toggle-status", headers=_hdr(token))
    assert r.status_code == 200 and r.json()["data"]["status"] == 0
    assert ctx.scalar(select(User).where(User.username == "sup_alpha")).status == 0
    r = ctx.client.post(
        "/api/auth/login", json={"username": "sup_alpha", "password": "Sup@2026x"}
    )
    assert r.status_code == 401 and r.json()["code"] == 40100
